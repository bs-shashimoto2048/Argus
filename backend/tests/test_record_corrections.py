"""読取値(正式値)の手動修正と、修正の監査履歴(reading_record_corrections)のテスト。一時DBのみ使用する(実運用DBは使わない)。"""
from __future__ import annotations

import io
import uuid
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import delete, select

from app.core.database import SessionLocal
from app.main import app
from app.models import InferenceSettings, LatestResult, Monitor, ReadingBaselineEvent, ReadingRecord, ReadingRecordCorrection, VideoSource
from app.services import reading_baseline_service as baseline_service
from app.services import reading_correction_service as corrections
from app.services.csv_export_service import JST
from runtime.hourly_record_worker import hourly_record_worker

pytestmark = pytest.mark.integration


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(hourly_record_worker, "start", lambda: None)
    with TestClient(app) as c:
        yield c


@pytest.fixture
def db():
    session = SessionLocal()
    session.info["monitors"] = []
    yield session
    session.rollback()
    for monitor_id in session.info["monitors"]:
        session.execute(delete(ReadingRecordCorrection).where(ReadingRecordCorrection.monitor_id == monitor_id))
        session.execute(delete(ReadingRecord).where(ReadingRecord.monitor_id == monitor_id))
        session.execute(delete(ReadingBaselineEvent).where(ReadingBaselineEvent.monitor_id == monitor_id))
        monitor = session.get(Monitor, monitor_id)
        if monitor is not None:
            session.delete(monitor)
    session.commit()
    session.close()


def make_monitor(db, expected_digits=7, decimal_position=None) -> Monitor:
    name = "t_corr_" + uuid.uuid4().hex[:8]
    monitor = Monitor(name=name, display_name=f"表示名{name[-4:]}", enabled=True, status="running", created_at=datetime(2026, 10, 1))
    monitor.inference = InferenceSettings(engine="cpp_onnx", model_id="digital.onnx", reading={"enabled": True, "expected_digits": expected_digits, "decimal_position": decimal_position})
    monitor.latest_result = LatestResult(value="215858", status="ok")
    monitor.source = VideoSource(source_type="url", url="http://example.invalid/x")
    db.add(monitor)
    db.commit()
    db.info["monitors"].append(monitor.id)
    return monitor


def add_record(db, monitor, hour, value, **over) -> ReadingRecord:
    bucket = datetime(2026, 10, 9, hour, tzinfo=JST)
    data = dict(monitor_id=monitor.id, monitor_name=monitor.display_name, hour_bucket=bucket.isoformat(), recorded_at=bucket.astimezone(timezone.utc).replace(tzinfo=None) + timedelta(seconds=15),
                value=value, numeric_value=value, raw_value="0" + value if value else None, raw_confidence=0.881, value_source="confirmed", confidence=0.93,
                validation_status="confirmed", display_status="normal", baseline_conflict=False, engine="cpp_onnx", model_id="digital.onnx",
                original_image_path=f"m/2026/10/09/{hour:02d}_original.jpg", overlay_image_path=f"m/2026/10/09/{hour:02d}_overlay.jpg", image_status="ok",
                inference_at=bucket.astimezone(timezone.utc).replace(tzinfo=None) + timedelta(seconds=14, milliseconds=243))
    data.update(over)
    record = ReadingRecord(**data)
    db.add(record)
    db.commit()
    return record


def carried_scenario(db):
    """07:00=215836(confirmed) / 08:00=215858(carried_forward。最新Rawは0215850で棄却) / 09:00=215865(confirmed)。"""
    monitor = make_monitor(db)
    r7 = add_record(db, monitor, 7, "215836")
    r8 = add_record(db, monitor, 8, "215858", value_source="carried_forward", validation_status="decrease_detected", raw_value="0215850", raw_confidence=0.881,
                    previous_value="215836", usage="22")
    r9 = add_record(db, monitor, 9, "215865", previous_value="215858", usage="7")
    return monitor, r7, r8, r9


def correct(client, record_id, **over):
    body = {"value": "215850", "reason": "実メーターの表示を確認", "operator": "tester", "rebase_current_baseline": False}
    body.update(over)
    return client.post(f"/api/records/{record_id}/correct", json=body)


# ---------- 修正できる記録・できない記録 ----------

def test_carried_forward_record_can_be_corrected(client, db):
    _monitor, _r7, r8, _r9 = carried_scenario(db)
    res = correct(client, r8.id)
    assert res.status_code == 200, res.text
    record = res.json()["record"]
    assert record["value"] == "215850" and record["is_corrected"] is True and record["correction_count"] == 1
    assert record["value_source"] == "carried_forward"  # 元の由来を消さない(manual_correctedへ上書きしない)
    assert record["original_value"] == "215858" and record["corrected_by"] == "tester" and record["corrected_at"].endswith("Z")


@pytest.mark.parametrize("over", [
    dict(baseline_conflict=True),
    dict(validation_status="decrease_detected"),
    dict(validation_status="rate_exceeded"),
])
def test_baseline_conflict_record_can_be_corrected(client, db, over):
    monitor = make_monitor(db)
    record = add_record(db, monitor, 8, "215858", **over)
    assert svc_correctable(record) is not None
    assert correct(client, record.id).status_code == 200


def svc_correctable(record):
    from app.services.reading_record_service import correctable_reason
    return correctable_reason(record)


def test_normal_confirmed_record_cannot_be_corrected(client, db):
    monitor = make_monitor(db)
    record = add_record(db, monitor, 8, "215858")  # confirmed / conflictなし
    res = correct(client, record.id)
    assert res.status_code == 409 and res.json()["detail"]["code"] == "NOT_CORRECTABLE"
    db.expire_all()
    assert db.get(ReadingRecord, record.id).value == "215858" and db.get(ReadingRecord, record.id).correction_count == 0
    assert db.scalars(select(ReadingRecordCorrection).where(ReadingRecordCorrection.record_id == record.id)).all() == []
    assert client.get(f"/api/records/{record.id}").json()["correctable"] is False  # UIに修正ボタンを出さない根拠


def test_record_without_a_value_cannot_be_corrected(client, db):
    monitor = make_monitor(db)
    record = add_record(db, monitor, 8, None, value_source="none", display_status="error", raw_value=None, numeric_value=None)
    assert correct(client, record.id).status_code == 409


def test_record_not_found(client):
    res = correct(client, 99999999)
    assert res.status_code == 404 and res.json()["detail"]["code"] == "RECORD_NOT_FOUND"


# ---------- 入力の検証 ----------

@pytest.mark.parametrize("over,code", [
    (dict(reason=""), "REASON_REQUIRED"), (dict(reason="   "), "REASON_REQUIRED"),
    (dict(operator=""), "OPERATOR_REQUIRED"), (dict(operator="  "), "OPERATOR_REQUIRED"),
    (dict(value="abc"), "INVALID_VALUE"), (dict(value=""), "INVALID_VALUE"), (dict(value="12-3"), "INVALID_VALUE"),
    (dict(value="21585000"), "INVALID_VALUE"),  # expected_digits(7桁)を超える
    (dict(value="215858"), "NO_CHANGE"),  # 現在の正式値と同じ
])
def test_invalid_requests_are_rejected_without_changes(client, db, over, code):
    _monitor, _r7, r8, _r9 = carried_scenario(db)
    res = correct(client, r8.id, **over)
    assert res.status_code == 422 and res.json()["detail"]["code"] == code, res.text
    db.expire_all()
    assert db.get(ReadingRecord, r8.id).value == "215858" and db.get(ReadingRecord, r8.id).usage == "22"
    assert db.scalars(select(ReadingRecordCorrection)).all() == [] or all(c.record_id != r8.id for c in db.scalars(select(ReadingRecordCorrection)).all())


def test_decimal_position_is_validated(client, db):
    monitor = make_monitor(db, expected_digits=7, decimal_position=1)
    record = add_record(db, monitor, 8, "372414.3", value_source="carried_forward", validation_status="decrease_detected", raw_value="372413.0")
    assert correct(client, record.id, value="372414").status_code == 422  # 小数部が1桁でない
    assert correct(client, record.id, value="372413.0").status_code == 200


# ---------- 監査履歴 ----------

def test_audit_row_records_old_new_reason_operator_and_evidence(client, db):
    monitor, _r7, r8, _r9 = carried_scenario(db)
    res = correct(client, r8.id, reason="  画像で0215850を確認  ", operator=" 山田 ")
    assert res.status_code == 200
    row = res.json()["correction"]
    assert (row["old_value"], row["new_value"], row["old_numeric_value"], row["new_numeric_value"]) == ("215858", "215850", "215858", "215850")
    assert (row["old_usage"], row["new_usage"]) == ("22", "14")
    assert (row["reason"], row["operator"]) == ("画像で0215850を確認", "山田")  # 前後の空白は除く
    assert (row["raw_value"], row["raw_confidence"], row["validation_status"], row["value_source"]) == ("0215850", 0.881, "decrease_detected", "carried_forward")
    assert row["overlay_image_path"].endswith("08_overlay.jpg") and row["original_image_path"].endswith("08_original.jpg")
    assert row["client_host"] and row["context"]["next_record"]["new_usage"] == "15" and row["context"]["rebase"] == {"requested": False, "performed": False}
    assert row["monitor_id"] == monitor.id and row["corrected_at"].endswith("Z")


def test_multiple_corrections_keep_the_full_history(client, db):
    _monitor, _r7, r8, _r9 = carried_scenario(db)
    assert correct(client, r8.id, value="215850", reason="1回目", operator="a").status_code == 200
    assert correct(client, r8.id, value="215852", reason="2回目", operator="b").status_code == 200
    assert correct(client, r8.id, value="215851", reason="3回目", operator="c").status_code == 200
    history = client.get(f"/api/records/{r8.id}/corrections").json()["corrections"]
    assert [(h["old_value"], h["new_value"], h["operator"], h["reason"]) for h in history] == [
        ("215852", "215851", "c", "3回目"), ("215850", "215852", "b", "2回目"), ("215858", "215850", "a", "1回目")]  # 新しい順・全件
    record = client.get(f"/api/records/{r8.id}").json()
    assert record["value"] == "215851" and record["correction_count"] == 3 and record["original_value"] == "215858"  # 元の正式値は最初の修正前のまま


def test_history_api_for_uncorrected_record_is_empty_and_404_for_unknown(client, db):
    _monitor, r7, _r8, _r9 = carried_scenario(db)
    assert client.get(f"/api/records/{r7.id}/corrections").json() == {"record_id": r7.id, "corrections": []}
    assert client.get("/api/records/99999999/corrections").status_code == 404


# ---------- 元証跡は変更しない ----------

def test_original_evidence_is_never_changed(client, db):
    _monitor, _r7, r8, _r9 = carried_scenario(db)
    before = {k: getattr(r8, k) for k in ("raw_value", "raw_confidence", "validation_status", "value_source", "original_image_path", "overlay_image_path", "inference_at", "recorded_at",
                                         "hour_bucket", "confidence", "display_status", "baseline_conflict", "engine", "model_id", "image_status")}
    assert correct(client, r8.id).status_code == 200
    assert correct(client, r8.id, value="215851", reason="再修正").status_code == 200
    db.expire_all()
    after = db.get(ReadingRecord, r8.id)
    assert {k: getattr(after, k) for k in before} == before
    assert after.value == "215851" and after.numeric_value == "215851"


# ---------- usage の再計算 ----------

def test_usage_is_recomputed_for_the_corrected_and_the_next_record(client, db):
    _monitor, _r7, r8, r9 = carried_scenario(db)
    res = correct(client, r8.id)  # 215858 -> 215850
    assert res.status_code == 200
    body = res.json()
    assert body["record"]["usage"] == "14"  # 215850 - 215836 (修正前は22)
    assert body["next_record"] == {"record_id": r9.id, "hour_bucket": r9.hour_bucket, "value": "215865", "value_source": "confirmed", "old_usage": "7", "new_usage": "15"}
    db.expire_all()
    nxt = db.get(ReadingRecord, r9.id)
    assert (nxt.value, nxt.previous_value, nxt.usage) == ("215865", "215850", "15")  # 次の記録の値は書き換えず、前回値とusageだけ整合させる(215865 - 215850)


def test_usage_null_conditions_are_kept(client, db):
    monitor = make_monitor(db)
    add_record(db, monitor, 7, "215836")
    # 手動修正済みは「運用者が確認した正式値」なので、記録時にdisplay_statusが異常系・baseline conflictだったとしても、usage計算上は信頼できる
    r8 = add_record(db, monitor, 8, "215858", value_source="carried_forward", display_status="read_error", usage=None, previous_value="215836")
    assert correct(client, r8.id, value="215850").status_code == 200
    db.expire_all()
    assert db.get(ReadingRecord, r8.id).usage == "14"
    m2 = make_monitor(db)
    add_record(db, m2, 7, "215836")
    c8 = add_record(db, m2, 8, "215858", baseline_conflict=True, validation_status="decrease_detected", usage=None, previous_value="215836")
    assert correct(client, c8.id, value="215850").status_code == 200
    db.expire_all()
    c8 = db.get(ReadingRecord, c8.id)
    assert c8.usage == "14" and c8.baseline_conflict is True  # 元のbaseline_conflictは証跡として残る
    # 値が前回より減る修正 → usageはnull(減少)
    m3 = make_monitor(db)
    add_record(db, m3, 7, "215836")
    d8 = add_record(db, m3, 8, "215858", value_source="carried_forward", validation_status="decrease_detected", usage="22", previous_value="215836")
    assert correct(client, d8.id, value="215800").status_code == 200
    db.expire_all()
    assert db.get(ReadingRecord, d8.id).usage is None


def test_next_record_usage_is_null_when_baseline_was_reset_in_between(client, db):
    monitor, _r7, r8, r9 = carried_scenario(db)
    db.add(ReadingBaselineEvent(monitor_id=monitor.id, monitor_name="m", occurred_at=r9.recorded_at - timedelta(minutes=30), action="rebase", old_value="1", new_value="2", reason="x", operator="o", client_host="h"))
    db.commit()
    assert correct(client, r8.id).status_code == 200
    db.expire_all()
    assert db.get(ReadingRecord, r9.id).usage is None  # reset/rebaseをまたぐ使用量はnull(従来の条件を維持)


def test_next_record_that_also_carries_the_same_value_is_not_rewritten(client, db):
    monitor, _r7, r8, _r9 = carried_scenario(db)
    r9 = db.scalars(select(ReadingRecord).where(ReadingRecord.monitor_id == monitor.id, ReadingRecord.hour_bucket.like("%T09:%"))).one()
    r9.value, r9.numeric_value, r9.value_source, r9.validation_status = "215858", "215858", "carried_forward", "decrease_detected"
    r9.usage = "0"
    db.commit()
    res = correct(client, r8.id)
    assert res.status_code == 200
    db.expire_all()
    nxt = db.get(ReadingRecord, r9.id)
    assert nxt.value == "215858" and nxt.correction_count == 0  # 同じ誤った値を保持していても、勝手に一括書き換えしない
    assert nxt.previous_value == "215850" and nxt.usage is None  # previous_valueだけ整合。未修正のcarried_forwardはusageを計算しない(8と自動計算しない)
    assert correct(client, r9.id, value="215851", reason="次の記録も個別に修正").status_code == 200  # 必要なら次のrecordも個別に修正できる


# ---------- baseline との関係 ----------

def test_baseline_is_not_changed_by_default(client, db, monkeypatch):
    called = []
    monkeypatch.setattr(baseline_service, "rebase_baseline", lambda *a, **k: called.append((a, k)))
    _monitor, _r7, r8, _r9 = carried_scenario(db)
    res = correct(client, r8.id)  # rebase_current_baseline 未指定(既定OFF)
    assert res.status_code == 200 and called == []
    assert res.json()["rebase"] == {"requested": False, "performed": False}


def test_baseline_is_rebased_only_when_explicitly_requested(client, db, monkeypatch):
    called = []
    monkeypatch.setattr(baseline_service, "rebase_baseline", lambda db_, monitor, value, reason, operator, host, force=False: called.append({"value": value, "reason": reason, "operator": operator, "force": force}))
    monkeypatch.setattr(baseline_service, "get_status", lambda db_, monitor: {"baseline": {"value": "215858"}})
    monitor, _r7, r8, _r9 = carried_scenario(db)
    res = correct(client, r8.id, rebase_current_baseline=True)
    assert res.status_code == 200, res.text
    assert len(called) == 1 and called[0]["value"] == "215850" and called[0]["operator"] == "tester" and called[0]["force"] is False
    assert "record" in called[0]["reason"] and "実メーター" in called[0]["reason"]
    assert res.json()["rebase"] == {"requested": True, "performed": True, "old_baseline": "215858", "new_baseline": "215850"}
    ctx = client.get(f"/api/records/{r8.id}/corrections").json()["corrections"][0]["context"]
    assert ctx["rebase"]["performed"] is True  # 修正の監査履歴にも、baseline再設定を実施したことが残る


def test_rebase_failure_does_not_change_the_record(client, db, monkeypatch):
    def failing(db_, monitor, value, reason, operator, host, force=False):
        raise baseline_service.ForceRequiredError("215900", value, baseline_service.Decimal("5"))

    monkeypatch.setattr(baseline_service, "rebase_baseline", failing)
    monkeypatch.setattr(baseline_service, "get_status", lambda db_, monitor: {"baseline": {"value": "215858"}})
    _monitor, _r7, r8, _r9 = carried_scenario(db)
    res = correct(client, r8.id, rebase_current_baseline=True)
    assert res.status_code == 409 and res.json()["detail"]["code"] == "REBASE_FORCE_REQUIRED"
    db.expire_all()
    assert db.get(ReadingRecord, r8.id).value == "215858" and db.get(ReadingRecord, r8.id).correction_count == 0  # 修正も行われない


def test_real_rebase_writes_both_audit_trails(client, db):
    monitor, _r7, r8, _r9 = carried_scenario(db)
    res = correct(client, r8.id, rebase_current_baseline=True)
    assert res.status_code == 200, res.text
    assert res.json()["rebase"]["performed"] is True
    events = baseline_service.list_events(db, monitor.id)
    assert events and events[0]["action"] == "rebase" and events[0]["new_value"] == "215850" and events[0]["operator"] == "tester"  # baseline側の監査履歴
    assert client.get(f"/api/records/{r8.id}/corrections").json()["corrections"][0]["new_value"] == "215850"  # 記録の修正側の監査履歴


# ---------- 表示・出力 ----------

def test_api_exposes_correction_state_and_evidence_fields(client, db):
    _monitor, _r7, r8, _r9 = carried_scenario(db)
    before = client.get(f"/api/records/{r8.id}").json()
    assert before["is_corrected"] is False and before["correctable"] is True and before["correctable_reason"] == "carried_forward"
    assert before["raw_confidence"] == 0.881 and before["confidence"] == 0.93 and before["snapshot_consistent"] is True
    correct(client, r8.id)
    listed = client.get("/api/records", params={"monitor_id": [r8.monitor_id], "limit": 10}).json()["items"]
    mine = next(i for i in listed if i["id"] == r8.id)
    assert mine["is_corrected"] is True and mine["value"] == "215850" and mine["original_value"] == "215858" and mine["raw_value"] == "0215850"


def test_excel_exports_the_corrected_value_and_keeps_the_evidence(client, db):
    monitor, _r7, r8, _r9 = carried_scenario(db)
    correct(client, r8.id)
    res = client.post("/api/records/export/excel", json={"monitor_ids": [monitor.id], "from": "2026-10-09T00:00:00+09:00", "to": "2026-10-10T00:00:00+09:00"})
    assert res.status_code == 200
    ws = load_workbook(io.BytesIO(res.content)).worksheets[0]
    rows = {r[0].value.hour: [c.value for c in r] for r in ws.iter_rows(min_row=2)}
    row = rows[8]
    assert row[3] == 215850 and row[5] == 14  # 確定値・使用量は修正後の値
    assert row[6] == "0215850" and row[9] == "carried_forward" and row[8] == "decrease_detected"  # 元証跡(Raw・判定・由来)は変わらない
    assert row[16] == 0.881  # Raw信頼度
    assert row[18] is True and row[19] == 1 and row[21] == 215858  # 修正済み / 修正回数 / 修正前値
    assert row[20] is not None  # 最終修正日時
    assert rows[9][5] == 15  # 次の記録のusageも再計算後の値
    assert rows[7][18] is False and rows[7][19] == 0 and rows[7][21] is None  # 未修正の記録


def test_correction_error_contract(db):
    with pytest.raises(corrections.CorrectionError) as info:
        corrections.correct_record(db, 99999999, "1", "r", "o")
    assert info.value.status == 404
