"""手動修正後のusageの順次再計算のテスト(usageは正式値同士の差。値の由来は問わない)。一時DBのみ使用する(実運用DBは使わない)。"""
from __future__ import annotations

import io
from datetime import timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import ReadingBaselineEvent, ReadingRecord
from app.services.reading_record_service import is_usage_eligible_record
from tests.test_record_corrections import add_record, client, correct, db, make_monitor  # noqa: F401  (fixtures)

pytestmark = pytest.mark.integration


def usages(db, monitor):
    db.expire_all()
    rows = db.scalars(select(ReadingRecord).where(ReadingRecord.monitor_id == monitor.id).order_by(ReadingRecord.hour_bucket)).all()
    return {int(r.hour_bucket[11:13]): r.usage for r in rows}


def cf(db, monitor, hour, value, **over):
    return add_record(db, monitor, hour, value, value_source="carried_forward", validation_status="decrease_detected", usage=None, **over)


def test_case1_carried_forward_corrected_gets_usage(client, db):
    m = make_monitor(db)
    add_record(db, m, 14, "100")
    r15 = cf(db, m, 15, "100")
    assert correct(client, r15.id, value="105").status_code == 200
    assert usages(db, m)[15] == "5"


def test_case2_next_confirmed_record_recovers(client, db):
    m = make_monitor(db)
    add_record(db, m, 14, "100")
    r15 = cf(db, m, 15, "100")
    add_record(db, m, 16, "110", usage=None)
    res = correct(client, r15.id, value="105")
    assert res.status_code == 200
    result = usages(db, m)
    assert (result[15], result[16]) == ("5", "5")
    body = res.json()
    assert body["next_record"]["record_id"] and body["recomputed_records"][0]["new_usage"] == "5"


def test_case3_carried_forward_after_a_correction_uses_the_formal_value_difference(client, db):
    m = make_monitor(db)
    add_record(db, m, 14, "100")
    r15 = cf(db, m, 15, "100")
    cf(db, m, 16, "105")
    add_record(db, m, 17, "112", usage=None)
    assert correct(client, r15.id, value="105").status_code == 200
    # 15:00=修正済み → usage 5。16:00=同じ値105を保持するcarried_forward → 0。17:00=新しいConfirmed値 → 105との差7
    result = usages(db, m)
    assert (result[15], result[16], result[17]) == ("5", "0", "7")


def test_case4_correcting_the_untrusted_record_recovers_in_order(client, db):
    m = make_monitor(db)
    add_record(db, m, 14, "100")
    r15 = cf(db, m, 15, "100")
    r16 = cf(db, m, 16, "105")
    add_record(db, m, 17, "112", usage=None)
    assert correct(client, r15.id, value="105").status_code == 200
    assert correct(client, r16.id, value="108").status_code == 200
    result = usages(db, m)
    assert (result[15], result[16], result[17]) == ("5", "3", "4")


def test_consecutive_corrected_records_recover_one_by_one(client, db):
    m = make_monitor(db)
    add_record(db, m, 15, "215862")
    r16 = cf(db, m, 16, "215862")
    r17 = cf(db, m, 17, "215862")
    add_record(db, m, 18, "215867", usage=None)
    assert correct(client, r16.id, value="215863").status_code == 200
    result = usages(db, m)
    assert (result[16], result[17], result[18]) == ("1", None, "5")  # 17は215862(修正した16の215863より減少)でnull。18は17の正式値215862との差5
    assert correct(client, r17.id, value="215864").status_code == 200
    result = usages(db, m)
    assert (result[16], result[17], result[18]) == ("1", "1", "3")


def test_case5_original_baseline_conflict_is_kept_as_evidence_and_does_not_affect_usage(client, db):
    m = make_monitor(db)
    add_record(db, m, 14, "100")
    r15 = add_record(db, m, 15, "100", baseline_conflict=True, value_source="carried_forward", validation_status="rate_exceeded", usage=None)
    assert is_usage_eligible_record(db.get(ReadingRecord, r15.id))  # baseline conflictはusage計算の可否に影響しない
    assert correct(client, r15.id, value="104").status_code == 200
    db.expire_all()
    rec = db.get(ReadingRecord, r15.id)
    assert (rec.baseline_conflict, rec.value_source, rec.validation_status, rec.raw_value) == (True, "carried_forward", "rate_exceeded", "0100")  # 元証跡は不変
    assert is_usage_eligible_record(rec) and rec.usage == "4"


def test_unmodified_carried_forward_is_eligible_when_the_video_and_reading_are_running(db):
    m = make_monitor(db)
    assert is_usage_eligible_record(cf(db, m, 15, "100"))
    assert not is_usage_eligible_record(cf(db, m, 16, "100", display_status="read_error"))  # 連続性を保証できない


def test_rebase_to_the_confirmed_value_does_not_break_continuity_but_other_resets_do(client, db):
    m = make_monitor(db)
    add_record(db, m, 14, "100")
    r15 = cf(db, m, 15, "100")
    # 修正値(105)への基準値再設定(運用者が確認した値) → 連続性を壊さない
    db.add(ReadingBaselineEvent(monitor_id=m.id, monitor_name="m", occurred_at=r15.recorded_at - timedelta(minutes=10), action="rebase", old_value="100", new_value="105", reason="確認", operator="o", client_host="h"))
    db.commit()
    assert correct(client, r15.id, value="105").status_code == 200
    assert usages(db, m)[15] == "5"
    # 無関係な値へのreset/rebaseは従来どおり使用量を無効にする
    m2 = make_monitor(db)
    add_record(db, m2, 14, "100")
    q15 = cf(db, m2, 15, "100")
    db.add(ReadingBaselineEvent(monitor_id=m2.id, monitor_name="m", occurred_at=q15.recorded_at - timedelta(minutes=10), action="reset", old_value="100", new_value=None, reason="x", operator="o", client_host="h"))
    db.commit()
    assert correct(client, q15.id, value="105").status_code == 200
    assert usages(db, m2)[15] == "0"  # 無関係な値へのreset後の最初の正常なrecordは、新しい区間の開始点(差分は計上しない)


def test_recompute_does_not_touch_unrelated_usage_or_evidence(client, db):
    m = make_monitor(db)
    add_record(db, m, 13, "90", usage="777")  # 無関係な過去のusageは書き換えない
    add_record(db, m, 14, "100")
    r15 = cf(db, m, 15, "100")
    r16 = add_record(db, m, 16, "110", usage=None)
    fields = ("value", "raw_value", "raw_confidence", "validation_status", "value_source", "original_image_path", "overlay_image_path")
    before = tuple(getattr(r16, f) for f in fields)
    assert correct(client, r15.id, value="105").status_code == 200
    db.expire_all()
    r16 = db.get(ReadingRecord, r16.id)
    assert tuple(getattr(r16, f) for f in fields) == before
    assert usages(db, m)[13] == "777"


def test_excel_exports_recomputed_usage_and_previous_value(client, db):
    m = make_monitor(db)
    add_record(db, m, 14, "100")
    r15 = cf(db, m, 15, "100")
    add_record(db, m, 16, "110", usage=None)
    assert correct(client, r15.id, value="105").status_code == 200
    res = client.post("/api/records/export/excel", json={"monitor_ids": [m.id], "from": "2026-10-09T00:00:00+09:00", "to": "2026-10-10T00:00:00+09:00"})
    assert res.status_code == 200
    ws = load_workbook(io.BytesIO(res.content)).worksheets[0]
    rows = {r[0].value.hour: [c.value for c in r] for r in ws.iter_rows(min_row=2)}
    assert rows[15][3] == 105 and rows[15][5] == 5 and rows[16][5] == 5
