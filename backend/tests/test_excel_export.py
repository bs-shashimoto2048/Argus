"""reading_recordsからのExcel出力(UI再設計 Phase 3)のテスト。一時DB・一時フォルダのみ使用する(実運用DB/実カメラは使わない)。"""
from __future__ import annotations

import io
import time
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import delete

from app.core.database import SessionLocal
from app.main import app
from app.models import Monitor, ReadingRecord
from app.routers import records as records_router
from app.services import excel_export_service as excel
from app.services import storage_settings_service as storage
from app.services.csv_export_service import JST
from runtime.hourly_record_worker import hourly_record_worker

pytestmark = pytest.mark.integration

HEADER = ["計測日時", "Monitor ID", "Monitor名", "確定値", "前回値", "使用量", "Raw値", "信頼度", "validation_status", "value_source",
          "display_status", "baseline_conflict", "engine", "model_id", "元画像パス", "推論画像パス"]


def jst(hour: int, day: int = 8, minute: int = 0) -> datetime:
    return datetime(2026, 10, day, hour, minute, tzinfo=JST).astimezone(timezone.utc).replace(tzinfo=None)


@pytest.fixture
def db():
    session = SessionLocal()
    session.info["monitors"] = []
    yield session
    session.rollback()
    for monitor_id in session.info["monitors"]:
        session.execute(delete(ReadingRecord).where(ReadingRecord.monitor_id == monitor_id))
        monitor = session.get(Monitor, monitor_id)
        if monitor is not None:
            session.delete(monitor)
    storage.update(session, image_root_folder="", excel_output_folder="")
    session.commit()
    session.close()


@pytest.fixture
def client(monkeypatch, tmp_path):
    monkeypatch.setattr(hourly_record_worker, "start", lambda: None)
    monkeypatch.setattr(storage, "default_excel_root", lambda: tmp_path / "default_exports")
    with TestClient(app) as c:
        yield c


def make_monitor(db, display_name="ガスメーター") -> Monitor:
    monitor = Monitor(name="t_xl_" + uuid.uuid4().hex[:8], display_name=display_name, created_at=datetime(2026, 10, 1))
    db.add(monitor)
    db.commit()
    db.info["monitors"].append(monitor.id)
    return monitor


def add_record(db, monitor, hour=18, day=8, **fields) -> ReadingRecord:
    data = dict(monitor_id=monitor.id, monitor_name=monitor.display_name, hour_bucket=f"2026-10-{day:02d}T{hour:02d}:00:00+09:00",
                recorded_at=jst(hour, day, 0), value="265803", numeric_value="265803", raw_value="0265803", value_source="confirmed",
                confidence=0.987, validation_status="confirmed", display_status="normal", baseline_conflict=False,
                engine="cpp_onnx", model_id="digital.onnx", image_status="not_saved")
    data.update(fields)
    record = ReadingRecord(**data)
    db.add(record)
    db.commit()
    return record


def export(client, ids, **extra):
    return client.post("/api/records/export/excel", json={"monitor_ids": ids, **extra})


def open_book(response):
    assert response.status_code == 200, response.text
    return load_workbook(io.BytesIO(response.content))


def rows_of(ws):
    return list(ws.iter_rows(min_row=2, values_only=True))


# --- ブック構造 ---

def test_one_sheet_per_monitor_and_multiple_monitors(client, db):
    a, b, c = make_monitor(db, "Aメーター"), make_monitor(db, "Bメーター"), make_monitor(db, "Cメーター")
    for m in (a, b, c):
        add_record(db, m, 18)
        add_record(db, m, 19, previous_value="265803", usage="4")
    book = open_book(export(client, [a.id, c.id], **{"from": "2026-10-08T00:00:00+09:00"}))
    assert book.sheetnames == ["Aメーター", "Cメーター"]  # 1 Monitor 1 worksheet(Bは含まない)
    for ws in book.worksheets:
        assert [cell.value for cell in ws[1]] == HEADER
        assert len(rows_of(ws)) == 2  # 18:00と19:00が古い順に並ぶ
        assert ws.freeze_panes == "A2" and ws.auto_filter.ref


def test_all_monitors_when_monitor_ids_is_empty(client, db):
    a, b = make_monitor(db, "全A"), make_monitor(db, "全B")
    add_record(db, a), add_record(db, b)
    names = open_book(export(client, [])).sheetnames
    assert "全A" in names and "全B" in names


def test_single_monitor(client, db):
    m = make_monitor(db, "単独")
    add_record(db, m)
    assert open_book(export(client, [m.id])).sheetnames == ["単独"]


# --- シート名 ---

def test_sheet_name_sanitize_and_length():
    assert excel.sanitize_sheet_name("a:b\\c/d?e*f[g]h") == "a_b_c_d_e_f_g_h"
    assert excel.sanitize_sheet_name("x" * 50) == "x" * 31
    assert excel.sanitize_sheet_name("'quoted'") == "quoted"
    assert excel.sanitize_sheet_name("末尾にタブ	") == "末尾にタブ"  # 末尾の空白/制御文字は「_」にせず除く
    used: set[str] = set()
    assert excel.unique_sheet_name("", 7, used) == "Monitor_7"
    assert excel.unique_sheet_name("History", 8, used) == "Monitor_8"  # Excelの予約名は使わない


def test_sheet_name_duplicates_are_made_unique():
    used: set[str] = set()
    first = excel.unique_sheet_name("同じ名前", 2, used)
    second = excel.unique_sheet_name("同じ名前", 3, used)
    third = excel.unique_sheet_name("同じ名前".upper(), 4, used)
    assert first == "同じ名前" and second == "同じ名前_3" and third == "同じ名前_4"
    long_a = excel.unique_sheet_name("y" * 40, 12, used)
    long_b = excel.unique_sheet_name("y" * 40, 13, used)
    assert len(long_a) == 31 and len(long_b) <= 31 and long_a != long_b and long_b.endswith("_13")  # 31文字超過でもIDを付けて回避
    assert excel.unique_sheet_name("ABC", 1, used) != excel.unique_sheet_name("abc", 2, used)  # 大文字小文字違いも別名にする


def test_duplicate_and_overlong_display_names_in_a_workbook(client, db):
    a, b = make_monitor(db, "重複:名[前]"), make_monitor(db, "重複_名_前_")
    c = make_monitor(db, "あ" * 40)
    for m in (a, b, c):
        add_record(db, m)
    book = open_book(export(client, [a.id, b.id, c.id]))
    assert len(book.sheetnames) == 3 and len({n.lower() for n in book.sheetnames}) == 3
    assert all(len(n) <= 31 and not set(n) & set(":\\/?*[]") for n in book.sheetnames)


# --- 期間 ---

def test_period_resolution():
    now = datetime(2026, 10, 9, 15, 30, tzinfo=JST)
    assert excel.resolve_period("today", None, None, now) == (datetime(2026, 10, 9, tzinfo=JST), datetime(2026, 10, 10, tzinfo=JST))
    assert excel.resolve_period("last_7_days", None, None, now) == (datetime(2026, 10, 3, tzinfo=JST), datetime(2026, 10, 10, tzinfo=JST))
    start, end = datetime(2026, 10, 1, tzinfo=JST), datetime(2026, 10, 8, tzinfo=JST)
    assert excel.resolve_period("custom", start, end) == (start, end)
    with pytest.raises(excel.ExportError):
        excel.resolve_period("custom", end, start)
    with pytest.raises(excel.ExportError):
        excel.resolve_period("yesterday", None, None)


def test_today_and_last_7_days_filter_records(client, db):
    m = make_monitor(db, "期間")
    today = datetime.now(timezone.utc).astimezone(JST).replace(hour=1, minute=0, second=0, microsecond=0)
    for days_ago in (0, 3, 8):
        when = today - timedelta(days=days_ago)
        add_record(db, m, hour=1, day=when.day, hour_bucket=when.isoformat(), recorded_at=when.astimezone(timezone.utc).replace(tzinfo=None))
    db.query(Monitor).filter(Monitor.id == m.id).update({"created_at": datetime(2020, 1, 1)})
    db.commit()
    assert len(rows_of(open_book(export(client, [m.id], period="today"))["期間"])) == 1
    assert len(rows_of(open_book(export(client, [m.id], period="last_7_days"))["期間"])) == 2  # 8日前は含まない


def test_custom_period_and_monitor_filter(client, db):
    a, b = make_monitor(db, "期間A"), make_monitor(db, "期間B")
    for day in (6, 7, 8, 9):
        add_record(db, a, 12, day), add_record(db, b, 12, day)
    book = open_book(export(client, [a.id], **{"from": "2026-10-07T00:00:00+09:00", "to": "2026-10-09T00:00:00+09:00"}))
    assert book.sheetnames == ["期間A"] and len(rows_of(book["期間A"])) == 2  # 7日・8日(toは含まない)


def test_invalid_period_and_from_after_to(client, db):
    m = make_monitor(db)
    add_record(db, m)
    assert export(client, [m.id], **{"from": "not-a-date"}).status_code == 422
    assert export(client, [m.id], **{"from": "2026-10-09T00:00:00+09:00", "to": "2026-10-08T00:00:00+09:00"}).status_code == 422
    same = "2026-10-08T00:00:00+09:00"
    assert export(client, [m.id], **{"from": same, "to": same}).status_code == 422
    assert export(client, [m.id], period="yesterday").status_code == 422


def test_empty_result_is_404_with_code(client, db):
    m = make_monitor(db, "空")
    res = export(client, [m.id])
    assert res.status_code == 404 and res.json()["detail"]["code"] == "NO_RECORDS"
    res = export(client, [m.id], save_to_server=True)
    assert res.status_code == 404


# --- セルの型と値 ---

def test_cell_types_and_values(client, db):
    m = make_monitor(db, "セル")
    add_record(db, m, 18, value="265803", raw_value="0265803", previous_value=None, usage=None)
    add_record(db, m, 19, value="265807", previous_value="265803", usage="4", raw_value="0265807", confidence=0.5)
    add_record(db, m, 20, value="372414.3", previous_value="372413.8", usage="0.5", raw_value="372414.3")
    ws = open_book(export(client, [m.id]))["セル"]
    r18, r19, r20 = rows_of(ws)
    assert isinstance(r19[3], (int, float)) and r19[3] == 265807  # 確定値は数値
    assert isinstance(r19[4], (int, float)) and r19[4] == 265803
    assert isinstance(r19[5], (int, float)) and r19[5] == 4
    assert r20[3] == 372414.3 and r20[5] == 0.5
    assert isinstance(r19[7], float) and r19[7] == 0.5  # 信頼度
    assert r18[6] == "0265803" and isinstance(r18[6], str)  # Rawは文字列で先頭0を保持
    assert ws.cell(2, 7).data_type == "s" and ws.cell(2, 4).data_type == "n"
    assert r18[4] is None and r18[5] is None  # nullは空セル
    assert ws.cell(2, 8).number_format == "0.000"


def test_usage_and_values_are_not_recomputed(client, db):
    m = make_monitor(db, "再計算なし")
    add_record(db, m, 18, value="100", previous_value="90", usage="999")  # DBの値が不整合でも、そのまま出力する
    row = rows_of(open_book(export(client, [m.id]))["再計算なし"])[0]
    assert row[5] == 999


def test_monitor_name_cell_is_normalized_but_db_is_untouched(client, db):
    assert excel.display_monitor_name("Drum Meter\t") == "Drum Meter"
    assert excel.display_monitor_name("\r\n Drum  Meter \t\n") == "Drum  Meter"  # 内部のスペースは維持
    assert excel.display_monitor_name(None) == ""
    m = make_monitor(db, "名前")
    record = add_record(db, m, 18, monitor_name="Drum Meter\t")
    ws = open_book(export(client, [m.id]))["名前"]
    assert ws.cell(2, 3).value == "Drum Meter"
    db.expire_all()
    assert db.get(ReadingRecord, record.id).monitor_name == "Drum Meter\t"  # 証跡はそのまま


def test_carried_forward_keeps_earlier_confirmed_value_with_rejected_raw(client, db):
    # 07:00の正式値215836。その後に正常確定した215858を、08:00の瞬間のRaw(0215850, decrease_detected)を棄却したうえで保持する。
    m = make_monitor(db, "保持")
    add_record(db, m, 7, value="215836", raw_value="0215836")
    add_record(db, m, 8, value="215858", numeric_value="215858", raw_value="0215850", previous_value="215836", usage="22",
               value_source="carried_forward", validation_status="decrease_detected")
    row = rows_of(open_book(export(client, [m.id]))["保持"])[1]
    assert (row[3], row[4], row[5], row[6], row[8], row[9]) == (215858, 215836, 22, "0215850", "decrease_detected", "carried_forward")


def test_datetime_is_jst_and_formatted(client, db):
    m = make_monitor(db, "日時")
    add_record(db, m, 18, recorded_at=jst(18) + timedelta(seconds=5))
    ws = open_book(export(client, [m.id]))["日時"]
    cell = ws.cell(2, 1)
    assert cell.value == datetime(2026, 10, 8, 18, 0, 5)  # recorded_at(UTC)をJSTへ変換
    assert cell.number_format == "yyyy/mm/dd hh:mm:ss"


def test_status_columns_and_carried_forward_and_null_row(client, db):
    m = make_monitor(db, "状態")
    add_record(db, m, 18, value="265803", value_source="carried_forward", validation_status="rejected_not_monotonic",
               display_status="read_error", baseline_conflict=True)
    add_record(db, m, 19, value=None, numeric_value=None, raw_value=None, value_source="none", confidence=None,
               validation_status=None, display_status="stopped", engine=None, model_id=None)
    first, second = rows_of(open_book(export(client, [m.id]))["状態"])
    assert first[8:12] == ("rejected_not_monotonic", "carried_forward", "read_error", True)
    assert second[3] is None and second[6] is None and second[7] is None and second[9] == "none"
    assert second[12] is None and second[13] is None


# --- 画像 ---

def test_image_paths_hyperlink_when_file_exists_and_text_when_missing(client, db, tmp_path):
    root = tmp_path / "img"
    (root / "m" / "2026").mkdir(parents=True)
    (root / "m" / "2026" / "a_original.jpg").write_bytes(b"x")
    storage.update(db, image_root_folder=str(root))
    m = make_monitor(db, "画像")
    add_record(db, m, 19, original_image_path="m/2026/a_original.jpg", overlay_image_path="m/2026/missing_overlay.jpg", image_status="ok")
    add_record(db, m, 20)  # 画像なし(空セル)
    ws = open_book(export(client, [m.id]))["画像"]
    row = rows_of(ws)
    expected = str(root / "m/2026/a_original.jpg")
    assert row[0][14] == expected  # 相対パス + 現在のimage_root
    assert ws.cell(2, 15).hyperlink is not None and Path(ws.cell(2, 15).hyperlink.target.replace("file:///", "")).name == "a_original.jpg"
    assert row[0][15] == str(root / "m/2026/missing_overlay.jpg") and ws.cell(2, 16).hyperlink is None  # 欠損: Excelは失敗させず文字列のみ
    assert row[1][14] is None and row[1][15] is None


def test_image_root_unreachable_still_exports_with_path_text(client, db, tmp_path):
    storage.update(db, image_root_folder=str(tmp_path / "no_such_root"))
    m = make_monitor(db, "不通")
    add_record(db, m, 19, original_image_path="a/b.jpg")
    ws = open_book(export(client, [m.id]))["不通"]
    assert ws.cell(2, 15).value == str(tmp_path / "no_such_root" / "a/b.jpg") and ws.cell(2, 15).hyperlink is None


def test_image_path_outside_root_is_never_linked():
    linker = excel.ImageLinker(Path("C:/images"))
    assert linker.resolve("../secret.jpg") == ("../secret.jpg", False)
    assert linker.resolve("C:\\Windows\\x.jpg")[1] is False
    assert linker.resolve(None) == (None, False)


# --- 保存 / ダウンロード ---

def test_browser_download_response(client, db):
    m = make_monitor(db, "DL")
    add_record(db, m, 18)
    res = export(client, [m.id], **{"from": "2026-10-01T00:00:00+09:00", "to": "2026-10-09T00:00:00+09:00"})
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert 'filename="Argus_MeterRecords_20261001_20261008.xlsx"' in res.headers["content-disposition"]
    assert res.content[:2] == b"PK"  # xlsx(zip)


def test_save_to_server_local_folder(client, db, tmp_path):
    out = tmp_path / "out"
    out.mkdir()
    storage.update(db, excel_output_folder=str(out))
    m = make_monitor(db, "保存")
    add_record(db, m, 18)
    res = export(client, [m.id], save_to_server=True, **{"from": "2026-10-08T00:00:00+09:00", "to": "2026-10-09T00:00:00+09:00"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["saved"] is True and body["total_rows"] == 1 and body["sheets"] == [{"monitor_id": m.id, "sheet_name": "保存", "rows": 1}]
    saved = Path(body["path"])
    assert saved.parent == out and saved.name == "Argus_MeterRecords_20261008.xlsx"
    assert load_workbook(saved).sheetnames == ["保存"]  # 保存したファイルが実際に開ける
    assert [p.name for p in out.iterdir()] == [saved.name]  # 一時ファイルは残らない


def test_save_to_server_uses_default_folder_when_unset(client, db, tmp_path):
    m = make_monitor(db, "既定")
    add_record(db, m, 18)
    body = export(client, [m.id], save_to_server=True).json()
    assert Path(body["folder"]) == tmp_path / "default_exports" and Path(body["path"]).is_file()


def test_filename_collision_is_not_overwritten(client, db, tmp_path):
    out = tmp_path / "out"
    out.mkdir()
    storage.update(db, excel_output_folder=str(out))
    m = make_monitor(db, "衝突")
    add_record(db, m, 18)
    (out / "Argus_MeterRecords_20261008.xlsx").write_bytes(b"existing")
    kwargs = dict(save_to_server=True, **{"from": "2026-10-08T00:00:00+09:00", "to": "2026-10-09T00:00:00+09:00"})
    first = export(client, [m.id], **kwargs).json()
    second = export(client, [m.id], **kwargs).json()
    assert first["filename"] == "Argus_MeterRecords_20261008_2.xlsx" and second["filename"] == "Argus_MeterRecords_20261008_3.xlsx"
    assert (out / "Argus_MeterRecords_20261008.xlsx").read_bytes() == b"existing"


def test_filenames():
    d = lambda y, mo, da: datetime(y, mo, da, tzinfo=JST)
    assert excel.export_filename(d(2026, 10, 1), d(2026, 10, 9)) == "Argus_MeterRecords_20261001_20261008.xlsx"
    assert excel.export_filename(d(2026, 10, 9), d(2026, 10, 10)) == "Argus_MeterRecords_20261009.xlsx"  # 今日のみ
    assert excel.export_filename(None, None) == "Argus_MeterRecords_All.xlsx"
    assert excel.safe_filename('a:b*c?.xlsx') == "a_b_c_.xlsx"


# --- 保存先の障害 ---

def test_save_failures_return_503_and_nothing_else(client, db, tmp_path):
    m = make_monitor(db, "障害")
    add_record(db, m, 18)
    storage.update(db, excel_output_folder=str(tmp_path / "missing_share"))  # UNC不通を模した、存在しない保存先
    res = export(client, [m.id], save_to_server=True)
    assert res.status_code == 503 and res.json()["detail"]["code"] == "EXPORT_SAVE_FAILED"
    # 失敗後も、ダウンロードや他のAPIは通常どおり使える
    assert export(client, [m.id]).status_code == 200
    assert client.get("/api/records/status").status_code == 200


def test_unc_share_that_does_not_respond_times_out(client, db, tmp_path, monkeypatch):
    m = make_monitor(db, "UNC")
    add_record(db, m, 18)
    storage.update(db, excel_output_folder=r"\\unreachable-host\share\argus")
    monkeypatch.setattr(excel, "SAVE_TIMEOUT_SECONDS", 0.3)
    monkeypatch.setattr(Path, "exists", lambda self: time.sleep(2) or False)
    started = time.monotonic()
    res = export(client, [m.id], save_to_server=True)
    assert res.status_code == 503 and res.json()["detail"]["code"] == "EXPORT_SAVE_TIMEOUT" and time.monotonic() - started < 5


def test_failed_export_does_not_touch_runtime_workers(client, db, tmp_path):
    from runtime.record_writer import record_writer
    before = (record_writer.status()["counts"], record_writer.is_running())
    m = make_monitor(db, "影響なし")
    add_record(db, m, 18)
    storage.update(db, excel_output_folder=str(tmp_path / "nope"))
    assert export(client, [m.id], save_to_server=True).status_code == 503
    after = (record_writer.status()["counts"], record_writer.is_running())
    assert before == after and not record_writer.status()["circuit_open"]  # RecordWriter/画像保存の状態は変わらない


def test_concurrent_export_is_rejected(client, db):
    m = make_monitor(db, "同時")
    add_record(db, m, 18)
    assert records_router._export_lock.acquire(blocking=False)
    try:
        res = export(client, [m.id])
        assert res.status_code == 409 and res.json()["detail"]["code"] == "EXPORT_IN_PROGRESS"
    finally:
        records_router._export_lock.release()


# --- 大量データ ---

def test_chunked_reading_keeps_order_and_all_rows(client, db, monkeypatch):
    monkeypatch.setattr(excel, "CHUNK_SIZE", 3)
    m = make_monitor(db, "チャンク")
    for hour in range(0, 8):
        add_record(db, m, hour, value=str(1000 + hour), numeric_value=str(1000 + hour))
    rows = rows_of(open_book(export(client, [m.id]))["チャンク"])
    assert [r[3] for r in rows] == [1000 + h for h in range(8)]  # チャンク境界をまたいでも、欠落・重複なく古い順


def test_ignores_inference_results_table(client, db):
    m = make_monitor(db, "データ源")
    assert export(client, [m.id]).status_code == 404  # reading_recordsが無ければ、他のテーブルがあっても出力しない
